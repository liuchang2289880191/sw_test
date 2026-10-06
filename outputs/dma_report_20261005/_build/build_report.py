from pathlib import Path
import csv
import json
import statistics
import zipfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
SRC = ROOT/'dma_results_20261005_183213_10036'
OUT = Path(__file__).resolve().parents[1]
FIG = OUT/'图表'
with (SRC/'summary.csv').open() as f:
    D = {(r['mode'],int(r['active_pes']),int(r['offset']),int(r['bytes'])):r
         for r in csv.DictReader(f)}
FITS = json.loads((ROOT/'outputs/dma_analysis_20261005_183213_10036/verification_and_fits.json').read_text(encoding='utf-8'))
S = [2**n for n in range(3,17)]
MODES = ['get','put','iget','iput']
def bw(m,n,o,s): return float(D[m,n,o,s]['aggregate_bytes_per_cycle'])
def cy(m,n,o,s): return float(D[m,n,o,s]['cycles_per_op'])
doc = Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
sec.footer_distance=Cm(.7)

def fontstyle(style, size, east='宋体',bold=False):
    style.font.name='Arial'; style.font.size=Pt(size); style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)

for name,size in [('Normal',10.5),('Title',22),('Heading 1',15),('Heading 2',12),('Caption',9)]:
    fontstyle(doc.styles[name],size,'黑体' if name in ['Title','Heading 1','Heading 2'] else '宋体',name!='Normal' and name!='Caption')
    fmt=doc.styles[name].paragraph_format
    fmt.space_after=Pt(6)
    if name.startswith('Heading'):
        fmt.space_before=Pt(10);fmt.keep_with_next=True
doc.styles['Normal'].paragraph_format.line_spacing=1.16
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
for el in doc.styles.element.iter(qn('w:pBdr')):
    el.getparent().remove(el)
doc.core_properties.title='SW39000 DMA 微基准结果分析'
doc.core_properties.subject='DMA 带宽 扩展比 地址对齐和逐核耗时'
doc.core_properties.author=''

def p(text,style=None,bold=False):
    para=doc.add_paragraph(style=style)
    run=para.add_run(text)
    if bold: run.bold=True
    return para
def h(text): return doc.add_heading(text,level=1)
def h2(text): return doc.add_heading(text,level=2)
def page(): doc.add_page_break()
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for cell,w in zip(t.columns,widths):cell.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side); el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,text,w) in enumerate(zip(cells,row,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para=cell.paragraphs[0];para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
            run=para.add_run(str(text));run.font.size=Pt(9.5)
            if i==0:
                run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'24445E' if i==0 else ('F1F5F8' if i%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shade)
        cant=OxmlElement('w:cantSplit');t.rows[i]._tr.get_or_add_trPr().append(cant)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def figure(name,caption,width=17.2):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.keep_with_next=True
    run=para.add_run();run.add_picture(str(FIG/(name+'.png')),width=Cm(width))
    shape=doc.inline_shapes[-1]
    shape._inline.docPr.set('descr',caption)
    para=p(caption,'Caption');para.alignment=WD_ALIGN_PARAGRAPH.CENTER

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' 页')
for r in footer.runs:r.font.size=Pt(9)

p('SW39000 DMA 微基准结果分析','Title')
p('实验日期 2026年10月5日    作业 8419757    队列 q_share')
p('本次测试显示，大消息下的 DMA 聚合吞吐趋于饱和；最明显的性能边界是 128 B 对齐写入。64 核传输 128 B 时，主存地址偏移 4 B 会使 put 完成耗时增加到对齐时的约 12.69 倍。这个边界应进入后续性能模型。')
h('1  实验配置与数据复核')
p('本次节点 cpuinfo 报告 cpu revision 为 sw39000，从核 SPE 频率为 2250 MHz；编译器为 swgcc/1473。结果按该运行环境记录，尚不能直接迁移为 SW26010Pro 参数。')
p('正式扫描包含 get、put、iget、iput 四种模式，活跃从核数为 1 或 64，主存偏移为 0 或 4 B，消息为 8 B 至 64 KiB 的 14 个倍增规模，共 224 项；另有 8 项正确性检查。独立复核了 7540 条逐核记录和 112 组扩展比，所有报告的 errors 为 0，summary 与 raw 数据一致。')
h2('计时与完成条件')
p('每核仅保留一个未完成请求。8 次预热不计时；小于 1 KiB 重复 10000 次，其余重复 1000 次。屏障、初始化、spawn/join、校验和 CSV 输出在计时区外。传输始终复用同一主存槽与 LDM 缓冲。')
p('summary 中的 cycles/op 取活跃核中最大累计周期除以重复次数。聚合带宽为消息字节数 × 重复次数 × 活跃核数 ÷ 最大累计周期。它是有效载荷吞吐，不包括内部协议和额外分段流量。')
h2('单核小消息完成开销')
rows=[]
for m in MODES:
    rng=[s for s in S if s<=(128 if m in ['get','iget'] else 64)]
    v=[cy(m,1,0,s) for s in rng]
    rows.append([m,f'8 至 {max(rng)} B',f'{min(v):.2f} 至 {max(v):.2f}',f'{statistics.median(v)/2.25:.2f}'])
table(['模式','消息区间','cycles/op 范围','条件换算 ns'],rows,[2.0,3.4,7.1,4.9])
p('读取基线约 441 cycles，写入约 234 至 240 cycles。数值包含接口、循环和应答等待，不能视为纯硬件启动延迟。')
p('单位说明：原始周期和 B/cycle 为主要结果。若周期计数器按报告的 2.25 GHz 从核频率递增，则 ns = cycles ÷ 2.25，GB/s = B/cycle × 2.25。本文换算值均采用这一条件，尚未独立标定计数器频率。')

page();h('2  DMA 带宽曲线')
p('单核带宽随消息增大持续上升。64 核读取在 4 至 64 KiB 约为 20.50 至 21.34 B/cycle；64 核写入在 8 至 64 KiB 约为 15.89 至 16.30 B/cycle，均出现较平的吞吐区间。')
figure('图1_DMA带宽曲线','图 1  四种 DMA 模式在不同核数与主存偏移下的聚合带宽')
h2('64 KiB 对齐传输的代表结果')
rows=[[m,f'{bw(m,1,0,65536):.3f}',f'{bw(m,64,0,65536):.3f}',f'{bw(m,64,0,65536)/bw(m,1,0,65536):.3f}',f'{bw(m,64,0,65536)*2.25:.2f}'] for m in MODES]
table(['模式','单核\nB/cycle','64 核聚合\nB/cycle','64 与 1 核\n带宽比','64 核条件换算\nGB/s'],rows,[1.7,3.7,4.0,3.7,4.3])
p('按 2.25 GHz 条件换算，64 核读取约为 47.8 GB/s，写入约为 36.0 GB/s。这些数值适用于单核组、单未完成请求和固定槽重复访问，不能作为全芯片或理论 DRAM 峰值。')
p('读写两类接口的完成定义也有差异。供应方等待完成表示数据已从本核 LDM 取走，不能额外解释成每一次物理主存写入均已落地。')

page();h('3  并发扩展与共享吞吐')
p('扩展比使用相同模式、消息大小和偏移条件下的 64 核聚合带宽除以单核带宽。小消息可通过并发提高吞吐，但收益随消息变大明显缩小。')
figure('图2_DMA并发扩展比','图 2  64 核相对单核的聚合带宽比  理想线性扩展值为 64')
rows=[]
for s in [8,64,128,1024,4096,16384,32768,65536]:
    rows.append([f'{s} B' if s<1024 else f'{s//1024} KiB',f'{bw("get",64,0,s)/bw("get",1,0,s):.2f}',f'{bw("put",64,0,s)/bw("put",1,0,s):.2f}'])
table(['消息大小','get 对齐扩展比','put 对齐扩展比'],rows,[4.0,6.7,6.7])
p('8 B 对齐读取扩展约 24.55 倍；到 64 KiB 时仅为 1.35 倍。64 KiB 对齐写入约为 0.91 倍，即 64 核聚合带宽比单核低约 9%。这符合大消息受到共享吞吐和争用约束的表现。')
p('目前只有 1 核与 64 核两个负载端点，无法确定从多少活跃核开始饱和，也无法区分内存控制器、DMA 注入、互连和地址映射各自的贡献。')

page();h('4  主存对齐与 128 B 写入边界')
p('偏移 4 B 仍满足接口的 4 B 对齐约束，但破坏 128 B 主存对齐。写入的偏移代价远大于读取，尤其集中在小中消息。')
figure('图3_DMA主存对齐影响','图 3  偏移 4 B 与对齐的完成耗时比  大于 1 表示偏移更慢')
h2('64 核下的偏移对照')
rows=[]
for m,s in [('get',128),('get',4096),('get',65536),('put',64),('put',128),('put',256),('put',4096),('put',65536)]:
    factor=cy(m,64,4,s)/cy(m,64,0,s)
    rows.append([m,str(s),f'{bw(m,64,0,s):.3f}',f'{bw(m,64,4,s):.3f}',f'{factor:.2f}',f'{(1-1/factor)*100:.1f}%'])
table(['模式','消息 B','对齐\nB/cycle','偏移 4 B\nB/cycle','耗时比','带宽下降'],rows,[1.5,2.0,3.7,4.1,2.7,3.4])
p('64 核对齐 put 从 64 B 增到 128 B 时，最大平均完成耗时从 2602.50 降到 533.08 cycles/op，带宽增加约 9.76 倍。128 B 加 4 B 偏移后耗时为 6764.15 cycles/op，带宽下降 92.1%；iput 的结果也呈现同样边界。')
p('手册说明 DMA 主存访问按 128 B 请求拆分，不对界地址会增加分段。这与边界现象相符，但额外分段本身尚不足以解释约 12.7 倍耗时。部分块写入路径与争用等机制需要进一步区分，当前不能断言具体硬件原因。')

page();h('5  操作完成耗时与非阻塞接口')
figure('图4_DMA操作完成耗时','图 4  get 与 put 的完成耗时  纵轴采用对数刻度')
p('get 单核小消息耗时较平稳，64 核下随消息增长逐渐受共同吞吐约束。put 在 64 核、128 B 对齐处出现耗时下降，应按分段行为处理，不能对全尺寸范围使用同一条线性延迟曲线。')
h2('阻塞与非阻塞的匹配比较')
rows=[]
for n in [1,64]:
    for sync,async_ in [('get','iget'),('put','iput')]:
        ratios=[bw(async_,n,0,s)/bw(sync,n,0,s) for s in S]
        rows.append([str(n),f'{async_}/{sync}',f'{min(ratios):.4f} 至 {max(ratios):.4f}',f'{statistics.median(ratios):.4f}'])
table(['活跃核数','带宽比','14 种消息的比值范围','比值中位数'],rows,[2.9,3.0,7.0,4.5])
p('get 与 iget 的差异基本在约 1% 内。单核 iput 相对 put 通常快约 1% 至 2%；64 核下两者仍很接近。每次 iget/iput 发起后立即等待，使非阻塞接口仍按串行方式使用，因而当前结果不反映流水化或计算通信重叠的收益。')
h2('测量边界')
p('手册将获取方完成定义为数据已写入其 LDM，将供应方完成定义为数据已从其 LDM 取走。put/iput 计时衡量当前接口的供应方完成边界；主核的内容校验在 join 后进行。读写基线不能简单解释成 DRAM 读写物理延迟的比较。')
p('每个配置只有一个独立计时区间。循环中的 1000 或 10000 次操作提高平均值的稳定程度，但不等同于多次独立实验；这里不对约 1% 的差别作统计显著性判断。')

page();h('6  逐核耗时与空间差异')
p('64 核下的小消息耗时随核编号变化。下图采用 id = 8r + c 排列，展示 128 B 读取与写入在两种主存偏移下的逐核平均耗时。各子图采用独立色标，应按数值比较。')
figure('图5_DMA逐核耗时热图','图 5  64 核并发时的 128 B DMA 逐核平均耗时',width=16.6)
rows=[]
for m,s in [('get',8),('get',4096),('put',64),('put',128),('put',65536)]:
    with (SRC/'raw'/f'dma_{m}_64pe_{s}B_offset0.csv').open() as f:
        vals=[float(r['cycles_per_op']) for r in csv.DictReader(f) if r['pe']!='aggregate']
    rows.append([m,str(s),f'{min(vals):.2f}',f'{statistics.median(vals):.2f}',f'{max(vals):.2f}'])
table(['模式','消息 B','最小 cycles/op','中位数 cycles/op','最大 cycles/op'],rows,[1.7,2.3,4.3,4.7,4.4])
p('get 8 B 的最大耗时约为最小值的 1.89 倍，get 4 KiB 时则十分接近。写入有随列编号重复的带状差异，但固定主存槽步长同时随核编号变化，地址映射与核位置的影响尚未分离。这些 DMA 图不能作为 RMA 的 2×2 router 拓扑证据。')

page();h('7  初步建模参数与后续实验')
p('在主存对齐条件下拟合 T(S) = α + βS。T 为 cycles/op，S 为消息字节数。单核选 1 至 64 KiB 的 7 个点，64 核选 4 至 64 KiB 的 5 个点，分别拟合不同模式与负载。斜率对应的聚合带宽为活跃核数除以 β。')
rows=[]
for f in FITS['fits']:
    rows.append([f['mode'],str(f['active_pes']),f'{f["alpha_cycles"]:.2f}',f'{f["beta_cycles_per_byte"]:.6f}',f'{f["effective_aggregate_bytes_per_cycle"]:.2f}',f'{f["r_squared"]:.6f}'])
table(['模式','核数','α\ncycles','β\ncycles/B','斜率对应聚合\nB/cycle','R²'],rows,[1.9,1.4,3.0,4.0,4.4,2.7])
p('拟合描述当前区间的大消息趋势。读取单核截距约 490 cycles，略高于小消息平台；64 核 iput 的近零截距是区间拟合结果，不能解释成零物理延迟。模型应分别保存小消息开销、大消息吞吐以及对齐条件，写入尤其需要保留 128 B 边界。')
h2('后续测量建议')
for text in [
    '增加独立重复和运行顺序变化，记录资源共享情况；q_share 名称本身不能证明独占程度。',
    '细扫 64、96、124、128、132、192、252、256、260 B，以及 0、4、64、124 B 偏移，定位写入边界。',
    '增加 2、4、8、16、32 核负载；改变槽步长和槽与核的对应关系，区分位置与内存映射。',
    '用独立缓冲扫描未完成请求窗口与轮转主存工作集，测队列能力和不同访问范围；保存 cache/LDM 配置并标定计数器频率。']:
    para=p('• '+text);para.paragraph_format.space_after=Pt(4)
h2('数据与手册依据')
p('数据集 dma_results_20261005_183213_10036：summary.csv 与 scaling.csv；raw 和 smoke 原始记录；source 内本次源码快照；build_info.txt 与 cpuinfo.txt。图表均由该数据集计算，未执行新的神威测试。')
p('SACA编程指南 v0.62：PDF 第 59 页说明 DMA 的 128 B 主存请求拆分；第 60 页定义 DMA 完成条件；第 88 页定义从核周期计数器。',bold=False)
p('适用范围：本次 SW39000 环境报告下的单核组、单未完成请求、固定槽预热后访问。迁移到 SW26010Pro、其他 cache 配置或大轮转工作集前需重新验证参数。')

path=OUT/'SW39000_DMA微基准分析报告.docx'
doc.save(path)
(FIG/'图表说明.txt').write_text('DMA 测试日期 2026-10-05，作业 8419757。\n原始数据集 dma_results_20261005_183213_10036。\nPNG 为 300 dpi 图，SVG 为矢量图。\n带宽使用 B/cycle；GB/s = B/cycle × 实际计数器频率 GHz。\n热图按 id=8r+c 排列，各子图使用独立色标，不能据此判断 RMA 2×2 拓扑。\n',encoding='utf-8')
with zipfile.ZipFile(OUT/'DMA中文图表.zip','w',zipfile.ZIP_DEFLATED) as z:
    for f in sorted(FIG.iterdir()): z.write(f,'图表/'+f.name)
print(path)
